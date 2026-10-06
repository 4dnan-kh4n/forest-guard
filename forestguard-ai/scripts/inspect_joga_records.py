"""Read-only inventory of user-supplied Joga records; never edits source files."""
import hashlib
import json
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET
from openpyxl import load_workbook

SOURCE = Path(r'G:\Papa Office Forms\2017 - 18 NEW RDF + KUP\JOGA')
OUTPUT = Path(__file__).resolve().parents[1]/'reports'/'joga_records'
OUTPUT.mkdir(parents=True,exist_ok=True)
keyword = re.compile(r'joga|tksx|278|311|320|326|348|350|330|328|latitude|longitude|gps|coordinate|\bN\s*22|\bE\s*76|बीट|कक्ष|परिक्षेत्र|वनमंडल|अक्षांश|देशांतर',re.I)
records = []
for index, path in enumerate(sorted(SOURCE.rglob('*'))):
    if not path.is_file() or path.suffix.lower() not in {'.docx','.xlsx'} or path.name.startswith('~$'):
        continue
    record = {'path':str(path),'relative_path':str(path.relative_to(SOURCE)),
              'bytes':path.stat().st_size,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
              'matches':[],'fonts':[],'media':[]}
    try:
        with zipfile.ZipFile(path) as archive:
            for member in archive.namelist():
                if '/media/' in member and not member.endswith('/'):
                    target = OUTPUT/'embedded_media'/f'{index}_{Path(member).name}'
                    target.parent.mkdir(exist_ok=True)
                    target.write_bytes(archive.read(member))
                    record['media'].append({'member':member,'extracted_path':str(target)})
            if path.suffix.lower()=='.docx':
                ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
                for member in archive.namelist():
                    if member=='word/document.xml' or re.match(r'word/(header|footer)\d+\.xml',member):
                        tree=ET.fromstring(archive.read(member))
                        fonts={value for element in tree.findall('.//w:rFonts',ns) for value in element.attrib.values()}
                        record['fonts']=sorted(set(record['fonts'])|fonts)
                        for p,paragraph in enumerate(tree.findall('.//w:p',ns),1):
                            text=''.join(t.text or '' for t in paragraph.findall('.//w:t',ns))
                            if text.strip() and (keyword.search(text) or p<=12):
                                record['matches'].append({'part':member,'paragraph':p,'text':text})
            else:
                workbook=load_workbook(path,read_only=True,data_only=False)
                for sheet in workbook:
                    for row_number,row in enumerate(sheet.iter_rows(),1):
                        values=[f'{cell.coordinate}: {cell.value}' for cell in row if cell.value is not None]
                        text=' | '.join(values)
                        if values and (keyword.search(text) or row_number<=14):
                            record['matches'].append({'sheet':sheet.title,'row':row_number,'text':text})
                workbook.close()
    except Exception as error:
        record['error']=f'{type(error).__name__}: {error}'
    records.append(record)
(OUTPUT/'inventory.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
print('Files inspected:',len(records),'Embedded media:',sum(len(r['media']) for r in records))
for record in records:
    print('\nFILE:',record['relative_path'])
    if record.get('error'):print('ERROR:',record['error'])
    print('Matched passages:',len(record['matches']),'Embedded images:',len(record['media']))
